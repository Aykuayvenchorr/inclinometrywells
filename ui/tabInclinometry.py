from datetime import datetime

from openpyxl import Workbook

from qgis.PyQt.QtWidgets import (
    QFileDialog,
    QMessageBox,
    QTableWidgetItem,
)
from qgis.core import (
    QgsCoordinateReferenceSystem,
    QgsCoordinateTransform,
    QgsFeature,
    QgsGeometry,
    QgsLineString,
    QgsPoint,
    QgsPointXY,
    QgsProject,
    QgsVectorLayer,
)

from ..modules.excel_reader import ExcelReader
from ..modules.geodezy import Geodezy
from ..modules.inclinometry import Inclinometry
from ..settings.TableColumns import TableColumns


IncCol = TableColumns.inclinometry


class TabInclinometry:
    """Вкладка «Инклинометрия»."""

    # ------------------------------------------------------------------
    # Константы
    # ------------------------------------------------------------------

    AZIMUTH_TYPES = (
        "Магнитный",
        "Истинный",
        "Дирекционный угол",
    )

    DATE_FORMAT = "%Y-%m-%d"

    DEFAULT_MAGNETIC_ERROR = "0.5/1.0/1.5"
    DEFAULT_ZENITH_ERROR = "0.5"
    DEFAULT_AZIMUTH_ERROR = "0.5"

    WGS84_AUTHID = "EPSG:4326"
    PULKOVO42_AUTHID = "EPSG:4284"

    WELLBORE_ACTUAL_TYPE = 1
    WELLBORE_ACTUAL_NAME = "Фактический ствол"

    # ------------------------------------------------------------------
    # Инициализация
    # ------------------------------------------------------------------

    def __init__(self, dialog):
        self.tab = dialog

        # Состояние
        self.rows = 0
        self.currentDate = None
        self.crs = QgsCoordinateReferenceSystem()
        self.AzimuthType = 0

        self.north_p = 0.0
        self.east_p = 0.0
        self.alt = 0.0

        self.inclinometry = Inclinometry()
        self.excel = ExcelReader()

        self._connect_signals()
        self._configure_azimuth_type()

    # ------------------------------------------------------------------
    # Настройка интерфейса
    # ------------------------------------------------------------------

    def _connect_signals(self):
        """Подключает обработчики сигналов интерфейса."""

        self.tab.btnLoadInclinometry.clicked.connect(
            self.load_inclinometry
        )

        self.tab.cmbAzimuthType.currentIndexChanged.connect(
            self.onAzimuthTypeChanged
        )

        self.tab.btnCalculateInclinometry.clicked.connect(
            self.calculateInclinometryAndDeviation
        )

        self.tab.btnExportExcel.clicked.connect(
            self.export_inclinometry_to_excel
        )

        self.tab.btnCalculateInclinometry.setEnabled(False)
        self.tab.tabInclTargetGoBtn.setEnabled(False)

    def _configure_azimuth_type(self):
        """Настраивает список типов азимута."""

        self.tab.cmbAzimuthType.clear()
        self.tab.cmbAzimuthType.addItems(self.AZIMUTH_TYPES)

        # По умолчанию — магнитный азимут.
        self.tab.cmbAzimuthType.setCurrentIndex(0)
        self.AzimuthType = 0

    def _set_default_error_values(self):
        """Устанавливает значения погрешностей по умолчанию."""

        if self.tab.cmbAzimuthType.currentIndex() == 0:
            self.tab.txtErrMagneticAzimuth.setText(
                self.DEFAULT_MAGNETIC_ERROR
            )
            self.tab.txtErrZenith.setText(
                self.DEFAULT_ZENITH_ERROR
            )
            self.tab.txtErrAzimuth.setText(
                self.DEFAULT_AZIMUTH_ERROR
            )

    # ------------------------------------------------------------------
    # Работа с таблицей
    # ------------------------------------------------------------------

    def _table(self):
        """Возвращает таблицу инклинометрии."""

        return self.tab.tableInclinometry

    def _set_cell(self, row, column, value):
        """Записывает значение в ячейку таблицы."""

        self._table().setItem(
            row,
            column,
            QTableWidgetItem(str(value)),
        )

    def _get_cell_text(self, row, column):
        """Возвращает текст ячейки."""

        item = self._table().item(row, column)

        if item is None:
            return ""

        return item.text().strip()

    def _get_float(self, row, column):
        """Возвращает числовое значение из ячейки."""

        text = self._get_cell_text(row, column)

        if not text:
            raise ValueError(
                f"Пустое значение в строке {row + 1}, "
                f"столбце {column + 1}."
            )

        return float(text)

    # ------------------------------------------------------------------
    # Загрузка инклинометрии
    # ------------------------------------------------------------------

    def load_inclinometry(self):
        """Загрузка данных инклинометрии из Excel."""

        filename = self.excel.open_file()

        if not filename:
            return

        table = self._table()

        # Очищаем старые данные.
        table.clearContents()
        table.setRowCount(0)

        self.excel.data.clear()
        self.rows = 0
        self.currentDate = None

        # Читаем Excel.
        self.excel.read_inclinometry(filename)

        if not self.excel.data:
            self.tab.btnCalculateInclinometry.setEnabled(False)
            self.tab.tabInclTargetGoBtn.setEnabled(False)
            return

        # Запоминаем CRS расчёта.
        self.crs = (
            self.tab
            .mQgsProjectionSelectionWidgetWellHead
            .crs()
        )

        table.setRowCount(len(self.excel.data))

        table.setUpdatesEnabled(False)

        try:
            self._populate_inclinometry_table()
        finally:
            table.setUpdatesEnabled(True)

        self.tab.btnCalculateInclinometry.setEnabled(True)
        self.tab.tabInclTargetGoBtn.setEnabled(False)

    def _populate_inclinometry_table(self):
        """Заполняет таблицу данными из Excel."""

        table = self._table()

        for row, item in enumerate(self.excel.data):
            self._set_cell(
                row,
                IncCol["MD"],
                item["md"],
            )

            self._set_cell(
                row,
                IncCol["ZENITH"],
                item["zenith"],
            )

            self._set_cell(
                row,
                IncCol["AZIMUTH"],
                item["azimuth"],
            )

            self.rows = row + 1

            # Первая точка траектории.
            if row == 0:
                current_date = datetime.now().strftime(
                    self.DATE_FORMAT
                )

                self._set_cell(
                    row,
                    IncCol["DATE"],
                    current_date,
                )

                north_text = self.tab.txtWellHeadNorth.text()
                east_text = self.tab.txtWellHeadEast.text()
                rotor_text = self.tab.txtWellHeadRotor.text()

                self._set_cell(
                    row,
                    IncCol["NORTH"],
                    north_text,
                )

                self._set_cell(
                    row,
                    IncCol["EAST"],
                    east_text,
                )

                self._set_cell(
                    row,
                    IncCol["TVDSS"],
                    rotor_text,
                )

                self.north_p = float(north_text)
                self.east_p = float(east_text)
                self.alt = float(rotor_text)

                self._set_default_error_values()

    # ------------------------------------------------------------------
    # Тип азимута
    # ------------------------------------------------------------------

    def onAzimuthTypeChanged(self, index):
        """Обрабатывает изменение типа азимута."""

        self.AzimuthType = index

    # ------------------------------------------------------------------
    # Системы координат
    # ------------------------------------------------------------------

    def _transform_point(
        self,
        north: float,
        east: float,
        target_authid: str,
    ) -> QgsPointXY:
        """
        Переводит точку из CRS расчёта
        в указанную географическую систему координат.
        """

        target_crs = QgsCoordinateReferenceSystem(
            target_authid
        )

        transform = QgsCoordinateTransform(
            self.crs,
            target_crs,
            QgsProject.instance(),
        )

        source_point = QgsPointXY(east, north)

        return transform.transform(source_point)

    def getPointPulkovo42(
        self,
        north: float,
        east: float,
    ) -> QgsPointXY:
        """Преобразует точку в Пулково-42."""

        return self._transform_point(
            north,
            east,
            self.PULKOVO42_AUTHID,
        )

    def getPointWGS84(
        self,
        north: float,
        east: float,
    ) -> QgsPointXY:
        """Преобразует точку в WGS84."""

        return self._transform_point(
            north,
            east,
            self.WGS84_AUTHID,
        )

    def formatTargetCoordinates(self, point):
        """Форматирует координаты цели в зависимости от CRS."""

        if self.crsCurrentTarget.isGeographic():
            north_text = f"{point.y():.10f}"
            east_text = f"{point.x():.10f}"
        else:
            north_text = f"{point.y():.3f}"
            east_text = f"{point.x():.3f}"

        return north_text, east_text

    # ------------------------------------------------------------------
    # Расчёт магнитного склонения
    # ------------------------------------------------------------------

    def _calculate_magnetic_declination(
        self,
        north: float,
        east: float,
        altitude: float,
        date: datetime,
    ) -> float:
        """
        Рассчитывает магнитное склонение.

        Используется только для магнитного азимута.
        """

        if self.AzimuthType != 0:
            return 0.0

        point_wgs84 = self.getPointWGS84(
            north,
            east,
        )

        return Geodezy.magnetic_declination(
            Geodezy.deg2rad(point_wgs84.y()),
            Geodezy.deg2rad(point_wgs84.x()),
            altitude / 1000.0,
            date,
            self.AzimuthType,
        )[0]

    # ------------------------------------------------------------------
    # Расчёт сближения меридианов
    # ------------------------------------------------------------------

    def _calculate_meridian_convergence(
        self,
        north: float,
        east: float,
    ) -> float:
        """
        Рассчитывает сближение меридианов.

        Используется для:
            0 — магнитного азимута;
            1 — истинного азимута.

        Для дирекционного угла возвращается 0.
        """

        if self.AzimuthType >= 2:
            return 0.0

        projection_type = Geodezy.getProjectionType(
            self.crs
        )

        central_meridian = Geodezy.getCentralMeridian_(
            self.crs
        )

        if projection_type == "utm":
            point_wgs84 = self.getPointWGS84(
                north,
                east,
            )

            return Geodezy.convergence_meridians(
                Geodezy.deg2rad(point_wgs84.y()),
                Geodezy.deg2rad(point_wgs84.x()),
                central_meridian,
                projection_type,
                self.AzimuthType,
            )

        if projection_type == "gauss_kruger":
            point_pulkovo42 = self.getPointPulkovo42(
                north,
                east,
            )

            return Geodezy.convergence_meridians(
                Geodezy.deg2rad(point_pulkovo42.y()),
                Geodezy.deg2rad(point_pulkovo42.x()),
                central_meridian,
                projection_type,
                self.AzimuthType,
            )

        return 0.0

    # ------------------------------------------------------------------
    # Расчёт траектории
    # ------------------------------------------------------------------

    def calculateInclinometry(
        self,
        north_col,
        east_col,
        tvdss_col,
        switch,
    ):
        """
        Рассчитывает основную траекторию скважины.

        Поправки на ошибки измерений здесь НЕ применяются.
        Они рассчитываются отдельно в calculateErrorPoints().

        Тип азимута:
            0 — магнитный;
            1 — истинный;
            2 — дирекционный угол.

        Формула итогового азимута:

            azimuth_grid =
                azimuth
                + magnetic_declination
                + gamma

        Для дирекционного угла:
            magnetic_declination = 0
            gamma = 0

        Для истинного:
            magnetic_declination = 0
            gamma = рассчитывается

        Для магнитного:
            magnetic_declination = рассчитывается
            gamma = рассчитывается
        """

        table = self._table()

        # Начальные координаты первой точки.
        north_p = self._get_float(0, IncCol["NORTH"])
        east_p = self._get_float(0, IncCol["EAST"])
        alt = self._get_float(0, IncCol["TVDSS"])

        depth_prev = 0.0
        zenith_prev = 0.0
        azimuth_grid_prev = 0.0

        self.currentDate = datetime.strptime(
            self._get_cell_text(0, IncCol["DATE"]),
            self.DATE_FORMAT,
        )

        for i in range(self.rows):
            depth = self._get_float(
                i,
                IncCol["MD"],
            )

            zenith = Geodezy.deg2rad(
                self._get_float(
                    i,
                    IncCol["ZENITH"],
                )
            )

            azimuth = Geodezy.deg2rad(
                self._get_float(
                    i,
                    IncCol["AZIMUTH"],
                )
            )

            dt = self.getRowDate(i)

            if i > 0:
                alt = self._get_float(
                    i - 1,
                    IncCol["TVDSS"],
                )

            # ----------------------------------------------------------
            # Магнитное склонение
            # ----------------------------------------------------------

            magnetic_declination = (
                self._calculate_magnetic_declination(
                    north_p,
                    east_p,
                    alt,
                    dt,
                )
            )

            # ----------------------------------------------------------
            # Сближение меридианов
            # ----------------------------------------------------------

            gamma = self._calculate_meridian_convergence(
                north_p,
                east_p,
            )

            # ----------------------------------------------------------
            # Итоговый азимут
            # ----------------------------------------------------------

            azimuth_grid = (
                azimuth
                + magnetic_declination
                + gamma
            )

            # ----------------------------------------------------------
            # Расчёт шага траектории
            # ----------------------------------------------------------

            if i > 0:
                dl = depth - depth_prev

                dNorth, dEast, dTVDSS = (
                    self.inclinometry.inclinometry_step(
                        dl,
                        azimuth_grid_prev,
                        azimuth_grid,
                        zenith_prev,
                        zenith,
                    )
                )

                if switch == 1:
                    self._set_cell(
                        i,
                        IncCol["DELTA_NORTH"],
                        f"{dNorth:.3f}",
                    )

                    self._set_cell(
                        i,
                        IncCol["DELTA_EAST"],
                        f"{dEast:.3f}",
                    )

                    self._set_cell(
                        i,
                        IncCol["DELTA_TVDSS"],
                        f"{dTVDSS:.3f}",
                    )

                north_p += dNorth
                east_p += dEast
                alt -= dTVDSS

                self._set_cell(
                    i,
                    north_col,
                    f"{north_p:.3f}",
                )

                self._set_cell(
                    i,
                    east_col,
                    f"{east_p:.3f}",
                )

                self._set_cell(
                    i,
                    tvdss_col,
                    f"{alt:.3f}",
                )

                depth_prev = depth
                zenith_prev = zenith
                azimuth_grid_prev = azimuth_grid

            # ----------------------------------------------------------
            # Запись расчётных параметров
            # ----------------------------------------------------------

            if switch == 1:
                self._set_cell(
                    i,
                    IncCol["DECLINATION"],
                    f"{Geodezy.rad2deg(magnetic_declination):.7f}",
                )

                self._set_cell(
                    i,
                    IncCol["CONVERGENCE"],
                    f"{Geodezy.rad2deg(gamma):.7f}",
                )

                self._set_cell(
                    i,
                    IncCol["GRID_AZIMUTH"],
                    f"{Geodezy.rad2deg(azimuth_grid) % 360.0:.7f}",
                )

        self.tab.tabInclTargetGoBtn.setEnabled(True)

    # ------------------------------------------------------------------
    # Расчёт погрешностей
    # ------------------------------------------------------------------

    def calculateErrorPoints(
        self,
        zenith_error: float,
        azimuth_error: float,
    ):
        """Рассчитывает точки эллипса неопределённости."""

        if self.rows < 2:
            return

        for i in range(1, self.rows):
            dt = self.getRowDate(i)

            magnetic_azimuth_error = Geodezy.deg2rad(
                self.currentErrorDeclination(dt)
            )

            north1 = self._get_float(
                i - 1,
                IncCol["NORTH"],
            )

            east1 = self._get_float(
                i - 1,
                IncCol["EAST"],
            )

            md1 = self._get_float(
                i - 1,
                IncCol["MD"],
            )

            north2 = self._get_float(
                i,
                IncCol["NORTH"],
            )

            east2 = self._get_float(
                i,
                IncCol["EAST"],
            )

            md2 = self._get_float(
                i,
                IncCol["MD"],
            )

            # В исходной логике l равен текущей MD.
            l = md2

            (
                a,
                b,
                points,
            ) = self.inclinometry.calculate_error_points(
                north1,
                east1,
                md1,
                north2,
                east2,
                md2,
                l,
                azimuth_error,
                zenith_error,
                magnetic_azimuth_error,
            )

            (
                north_left,
                east_left,
                tvdss_left,

                north_right,
                east_right,
                tvdss_right,

                north_up,
                east_up,
                tvdss_up,

                north_down,
                east_down,
                tvdss_down,
            ) = points

            # Основные параметры эллипса.
            self._set_cell(
                i,
                IncCol["a"],
                f"{a:.3f}",
            )

            self._set_cell(
                i,
                IncCol["b"],
                f"{b:.3f}",
            )

            # Точки эллипса.
            point_values = {
                "NORTH_LEFT": north_left,
                "EAST_LEFT": east_left,
                "TVDSS_LEFT": tvdss_left,

                "NORTH_RIGHT": north_right,
                "EAST_RIGHT": east_right,
                "TVDSS_RIGHT": tvdss_right,

                "NORTH_TOP": north_up,
                "EAST_TOP": east_up,
                "TVDSS_TOP": tvdss_up,

                "NORTH_DOWN": north_down,
                "EAST_DOWN": east_down,
                "TVDSS_DOWN": tvdss_down,
            }

            for column_name, value in point_values.items():
                self._set_cell(
                    i,
                    IncCol[column_name],
                    f"{value:.3f}",
                )

    # ------------------------------------------------------------------
    # Общий расчёт
    # ------------------------------------------------------------------

    def calculateInclinometryAndDeviation(self):
        """
        Рассчитывает траекторию скважины,
        фактический ствол и эллипс неопределённости.
        """

        calculation_crs = (
            self.tab
            .mQgsProjectionSelectionWidgetWellHead
            .crs()
        )

        if calculation_crs.isGeographic():
            QMessageBox.warning(
                self.tab,
                "Ошибка",
                "Система координат на вкладке "
                "Позиции/Устья должна быть прямоугольной.\n"
                "Также необходимо перезагрузить инклинометрию.",
            )
            return

        # Основная траектория.
        self.calculateInclinometry(
            north_col=IncCol["NORTH"],
            east_col=IncCol["EAST"],
            tvdss_col=IncCol["TVDSS"],
            switch=1,
        )

        # Фактический ствол.
        self.createWellbore()

        # Эллипс неопределённости.
        self.calculateErrorPoints(
            zenith_error=Geodezy.deg2rad(
                float(
                    self.tab.txtErrZenith.text()
                )
            ),
            azimuth_error=Geodezy.deg2rad(
                float(
                    self.tab.txtErrAzimuth.text()
                )
            ),
        )

    # ------------------------------------------------------------------
    # Работа с вкладкой целей
    # ------------------------------------------------------------------

    def targetTabActivate(self):
        """Активирует вкладку целей."""

        targets_index = self.tab.tabWidget.indexOf(
            self.tab.tabTargets
        )

        self.tab.tabWidget.setTabEnabled(
            targets_index,
            True,
        )

        self.tab.tabWidget.setCurrentWidget(
            self.tab.tabTargets
        )

    # ------------------------------------------------------------------
    # Дата строки
    # ------------------------------------------------------------------

    def currentErrorDeclination(
        self,
        dt: datetime,
    ) -> float:
        """
        Возвращает погрешность магнитного азимута
        в зависимости от даты.

        Формат:
            0.5/1.0/1.5

        где:
            до 2025 года     → первое значение;
            2025–2029       → второе значение;
            с 2030 года     → третье значение.
        """

        errors = (
            self.tab.txtErrMagneticAzimuth
            .text()
            .strip()
        )

        error_values = [
            float(value.strip())
            for value in errors.split("/")
        ]

        if dt < datetime(2025, 1, 1):
            return error_values[0]

        if dt < datetime(2030, 1, 1):
            return error_values[1]

        return error_values[2]

    def getRowDate(self, row: int) -> datetime:
        """
        Возвращает дату строки.

        Если дата в текущей строке отсутствует,
        используется последняя указанная выше дата.
        """

        table = self._table()

        for current_row in range(row, -1, -1):
            item = table.item(
                current_row,
                IncCol["DATE"],
            )

            if item is None:
                continue

            date_text = item.text().strip()

            if not date_text:
                continue

            self.currentDate = datetime.strptime(
                date_text,
                self.DATE_FORMAT,
            )

            break

        return self.currentDate

    # ------------------------------------------------------------------
    # Wellbore
    # ------------------------------------------------------------------

    def createWellbore(self):
        """
        Создаёт фактический ствол скважины
        в слое wellbore.

        Координаты:
            X = EAST
            Y = NORTH
            Z = TVDSS

        Геометрия:
            LineStringZ
        """

        if self.rows < 2:
            QMessageBox.warning(
                self.tab,
                "Внимание",
                "Для построения ствола необходимо минимум "
                "две точки инклинометрии.",
            )
            return False

        layer_wellbore = self._find_wellbore_layer()

        if layer_wellbore is None:
            QMessageBox.warning(
                self.tab,
                "Внимание",
                "Слой wellbore не найден в проекте.",
            )
            return False

        if not self.crs.isValid():
            QMessageBox.warning(
                self.tab,
                "Ошибка",
                "Не определена система координат расчёта.",
            )
            return False

        crs_layer = layer_wellbore.crs()

        if not crs_layer.isValid():
            QMessageBox.warning(
                self.tab,
                "Ошибка",
                "Не определена система координат слоя wellbore.",
            )
            return False

        geometry = self._build_wellbore_geometry()

        if geometry is None:
            return False

        # Преобразование CRS при необходимости.
        if self.crs != crs_layer:
            transform = QgsCoordinateTransform(
                self.crs,
                crs_layer,
                QgsProject.instance(),
            )

            try:
                geometry.transform(transform)
            except Exception as exc:
                QMessageBox.warning(
                    self.tab,
                    "Ошибка",
                    "Не удалось преобразовать геометрию "
                    f"ствола в CRS слоя:\n{exc}",
                )
                return False

        if not layer_wellbore.isEditable():
            if not layer_wellbore.startEditing():
                QMessageBox.warning(
                    self.tab,
                    "Ошибка",
                    "Не удалось перевести слой wellbore "
                    "в режим редактирования.",
                )
                return False

        feature = QgsFeature(
            layer_wellbore.fields()
        )

        feature.setGeometry(geometry)

        self._fill_wellbore_attributes(
            feature,
            layer_wellbore,
        )

        if not layer_wellbore.addFeature(feature):
            layer_wellbore.rollBack()

            QMessageBox.warning(
                self.tab,
                "Ошибка",
                "Не удалось добавить фактический ствол "
                "в слой wellbore.",
            )

            return False

        if not layer_wellbore.commitChanges():
            layer_wellbore.rollBack()

            QMessageBox.warning(
                self.tab,
                "Ошибка",
                "Не удалось сохранить фактический ствол "
                "в слой wellbore.",
            )

            return False

        layer_wellbore.triggerRepaint()

        self._set_layer_visible(
            layer_wellbore
        )

        return True

    def _find_wellbore_layer(self):
        """Находит слой wellbore в проекте."""

        layer_name = (
            self.tab
            .tabSettingsBoresMLCBox
            .currentText()
        )

        for layer in QgsProject.instance().mapLayers().values():
            if not isinstance(
                layer,
                QgsVectorLayer,
            ):
                continue

            if layer.name() == layer_name:
                return layer

        return None

    def _build_wellbore_geometry(self):
        """
        Создаёт LineStringZ из расчётных точек.
        """

        table = self._table()
        points = []

        for row in range(self.rows):
            try:
                north = self._get_float(
                    row,
                    IncCol["NORTH"],
                )

                east = self._get_float(
                    row,
                    IncCol["EAST"],
                )

                tvdss = self._get_float(
                    row,
                    IncCol["TVDSS"],
                )

            except (TypeError, ValueError):
                continue

            # X = EAST
            # Y = NORTH
            # Z = TVDSS
            points.append(
                QgsPoint(
                    east,
                    north,
                    tvdss,
                )
            )

        if len(points) < 2:
            QMessageBox.warning(
                self.tab,
                "Внимание",
                "Не удалось получить достаточное "
                "количество расчётных точек для "
                "построения ствола.",
            )
            return None

        # QgsPoint содержит Z, поэтому QgsLineString
        # создаёт LineStringZ.
        line = QgsLineString(points)

        return QgsGeometry(line)

    def _fill_wellbore_attributes(
        self,
        feature: QgsFeature,
        layer: QgsVectorLayer,
    ):
        """Заполняет атрибуты фактического ствола."""

        fields = layer.fields()

        if fields.indexOf("id") >= 0:
            feature["id"] = self._get_next_feature_id(
                layer
            )

        if fields.indexOf("type") >= 0:
            feature["type"] = self.WELLBORE_ACTUAL_TYPE

        if fields.indexOf("name") >= 0:
            feature["name"] = self.WELLBORE_ACTUAL_NAME

        if fields.indexOf("rel") >= 0:
            feature["rel"] = True

    def _get_next_feature_id(
        self,
        layer: QgsVectorLayer,
    ) -> int:
        """Возвращает следующий свободный числовой ID."""

        ids = []

        for feature in layer.getFeatures():
            value = feature["id"]

            if value is None:
                continue

            try:
                ids.append(int(value))
            except (TypeError, ValueError):
                continue

        return max(ids, default=0) + 1

    def _set_layer_visible(
        self,
        layer: QgsVectorLayer,
    ):
        """Включает отображение слоя в дереве проекта."""

        tree_layer = (
            QgsProject.instance()
            .layerTreeRoot()
            .findLayer(layer.id())
        )

        if tree_layer is not None:
            tree_layer.setItemVisibilityChecked(
                True
            )

    # ------------------------------------------------------------------
    # Экспорт Excel
    # ------------------------------------------------------------------

    def export_inclinometry_to_excel(self):
        """Экспортирует таблицу инклинометрии в Excel."""

        file_path, _ = QFileDialog.getSaveFileName(
            self.tab,
            "Сохранить инклинометрию",
            "",
            "Excel (*.xlsx)",
        )

        if not file_path:
            return

        if not file_path.lower().endswith(".xlsx"):
            file_path += ".xlsx"

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Инклинометрия"

        table = self._table()

        # Заголовки.
        for column in range(table.columnCount()):
            header = table.horizontalHeaderItem(
                column
            )

            if header is not None:
                worksheet.cell(
                    row=1,
                    column=column + 1,
                    value=header.text(),
                )

        # Данные.
        for row in range(table.rowCount()):
            for column in range(table.columnCount()):
                item = table.item(
                    row,
                    column,
                )

                value = (
                    item.text()
                    if item is not None
                    else ""
                )

                worksheet.cell(
                    row=row + 2,
                    column=column + 1,
                    value=value,
                )

        self._auto_fit_excel_columns(
            worksheet,
            table.columnCount(),
            table.rowCount(),
        )

        # Закрепляем заголовок.
        worksheet.freeze_panes = "A2"

        try:
            workbook.save(file_path)

        except Exception as exc:
            QMessageBox.critical(
                self.tab,
                "Ошибка",
                "Не удалось сохранить файл Excel:\n"
                f"{exc}",
            )

    @staticmethod
    def _auto_fit_excel_columns(
        worksheet,
        column_count: int,
        row_count: int,
    ):
        """Автоматически подбирает ширину столбцов Excel."""

        for column in range(
            1,
            column_count + 1,
        ):
            max_length = 0

            for row in range(
                1,
                row_count + 2,
            ):
                cell = worksheet.cell(
                    row=row,
                    column=column,
                )

                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            worksheet.column_dimensions[
                worksheet.cell(
                    row=1,
                    column=column,
                ).column_letter
            ].width = min(
                max_length + 2,
                40,
            )